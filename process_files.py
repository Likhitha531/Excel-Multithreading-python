
import os
import csv
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import load_workbook

folder = r"D:\Excel_Multithreading_50GB\excel_data"
output = r"D:\Excel_Multithreading_50GB\results\results_checkpoint.csv"

search_word = "Hello"
number_of_threads = 10
batch_size = 10

os.makedirs(os.path.dirname(output), exist_ok=True)

def process_file(path):
    try:
        workbook = load_workbook(path,read_only=True, data_only=True)

        record_count = 0
        hello_count = 0

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                record_count += 1

                for value in row:
                    if value is not None:
                        if search_word.lower() in str(value).lower():
                            hello_count += 1

        workbook.close()

        return [os.path.basename(path),record_count,hello_count, "Success"]

    except Exception as error:
        return [os.path.basename(path), 0, 0,"Error: " + str(error)]


files = sorted(
    os.path.join(folder, filename)
    for filename in os.listdir(folder)
    if filename.lower().endswith(".xlsx"))

completed_files = set()

if os.path.exists(output):
    with open(output, "r", newline="", encoding="utf-8") as file:
        reader = csv.reader(file)
        next(reader, None)

        for row in reader:
            if len(row) >= 4 and row[3] == "Success":
                completed_files.add(row[0])

files = [
    path for path in files
    if os.path.basename(path) not in completed_files]

print("Total Excel files:", len(completed_files) + len(files))
print("Already completed:", len(completed_files))
print("Files remaining:", len(files))
print("Threads:", number_of_threads)
print("Batch size:", batch_size)

start_time = time.time()
processed_count = 0

if not os.path.exists(output):
    with open(output, "w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow([
            "File",
            "Record Count",
            "Hello Count",
            "Status"])

for batch_start in range(0, len(files), batch_size):
    batch = files[batch_start:batch_start + batch_size]
    results = []

    with ThreadPoolExecutor(max_workers=number_of_threads) as executor:
        tasks = [executor.submit(process_file, path) for path in batch]

        for task in as_completed(tasks):
            result = task.result()
            results.append(result)

            print(
                result[0],
                "| Records:", result[1],
                "| Hello count:", result[2],
                "|", result[3])

    with open(output, "a", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerows(results)

    processed_count += len(results)
    print(
        "\nBatch completed:",
        processed_count,
        "of",
        len(files),
        "remaining files")

    elapsed = time.time() - start_time
    print("Elapsed time:", round(elapsed / 60, 2), "minutes\n")

print("Processing session completed.")
print("Files processed this session:", processed_count)
print("Total execution time:", round(time.time() - start_time, 2), "seconds")
print("Results saved to:", output)
